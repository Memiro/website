"""The workbook from the admin: the button on the card of a product and the action on the list of sections."""

from collections.abc import AsyncIterator
from http import HTTPStatus
from io import BytesIO

import pytest
from django.test import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import create_async_engine

from memiro.adapters.xlsx.pricing_workbook import CALCULATION, CHECK
from memiro.presentation.django_admin.workbook import ONE_AT_A_TIME, XLSX
from tests.common.factory.catalog import CATEGORY, PRODUCT, SECOND_CATEGORY
from tests.integration.prime import prime_second_category

pytestmark = pytest.mark.usefixtures("admin_site", "primed_catalog")

APP = "memiro"
PRODUCT_WORKBOOK_URL = f"/admin/{APP}/product/{PRODUCT}/workbook/"
SECTIONS_URL = f"/admin/{APP}/category/"
WORKBOOK_ACTION = "download_workbook"


@pytest.fixture
async def second_section(admin_database_url: str) -> AsyncIterator[None]:
    """Put one more section next to the mirrors, so the action has two rows to refuse."""
    engine = create_async_engine(admin_database_url)
    try:
        await prime_second_category(engine, name="Шкафы", slug="cabinets", sort_order=2, is_published=True)
        yield
    finally:
        await engine.dispose()


def _action(*sections: object) -> dict[str, object]:
    """Spell the action on the list of sections the way the changelist posts it."""
    return {
        "action": WORKBOOK_ACTION,
        "_selected_action": [str(section) for section in sections],
        "index": "0",
    }


async def test_the_owner_downloads_the_workbook_from_the_card_of_a_product(owner_client: AsyncClient) -> None:
    """The button hands over a spreadsheet named after the product, not a page about one."""
    response = await owner_client.get(PRODUCT_WORKBOOK_URL)

    assert response.status_code == HTTPStatus.OK
    assert response.headers["Content-Type"] == XLSX
    assert "zerkalo-v-rame.xlsx" in response.headers["Content-Disposition"]


async def test_the_downloaded_workbook_opens_as_a_book_of_the_product(owner_client: AsyncClient) -> None:
    """What arrives is the whole book: the calculator and the check sheet of the engine."""
    response = await owner_client.get(PRODUCT_WORKBOOK_URL)

    workbook = load_workbook(BytesIO(response.content))
    assert {CALCULATION, CHECK} <= set(workbook.sheetnames)


async def test_the_card_of_a_product_offers_the_button(owner_client: AsyncClient) -> None:
    """The owner finds the book where his prices are — under the card that shows them."""
    response = await owner_client.get(f"/admin/{APP}/product/{PRODUCT}/change/")

    assert PRODUCT_WORKBOOK_URL in response.content.decode()


async def test_the_owner_downloads_the_workbook_of_one_section(owner_client: AsyncClient) -> None:
    """A section is enough: the book of a dictionary needs no product at all."""
    response = await owner_client.post(SECTIONS_URL, _action(CATEGORY))

    assert response.status_code == HTTPStatus.OK
    assert response.headers["Content-Type"] == XLSX


async def test_the_action_refuses_two_sections_at_once(
    owner_client: AsyncClient,
    second_section: None,  # noqa: ARG001
) -> None:
    """Two sections are two books, and the owner is told so instead of getting one of them."""
    response = await owner_client.post(SECTIONS_URL, _action(CATEGORY, SECOND_CATEGORY), follow=True)

    assert ONE_AT_A_TIME in response.content.decode()


async def test_the_workbook_answers_nobody_who_may_not_change_the_product(clerk_client: AsyncClient) -> None:
    """Staff without the permission of the card is refused the book made out of it."""
    response = await clerk_client.get(PRODUCT_WORKBOOK_URL)

    assert response.status_code == HTTPStatus.FORBIDDEN


async def test_the_workbook_sits_behind_the_login_of_the_admin(admin_site: None) -> None:  # noqa: ARG001
    """An anonymous visitor is sent to the login page, not handed the owner's prices."""
    response = await AsyncClient().get(PRODUCT_WORKBOOK_URL)

    assert response.status_code == HTTPStatus.FOUND
    assert response.headers["location"].startswith("/admin/login/")
