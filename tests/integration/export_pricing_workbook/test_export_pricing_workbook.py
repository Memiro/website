"""The workbook out of the live catalogue: what the owner gets, and what he is refused."""

from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from dishka import AsyncContainer
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncEngine

from memiro.adapters.xlsx.pricing_workbook import CHECK, DATA_START
from memiro.application.errors.catalog import CategoryNotFoundError, ProductNotFoundError
from memiro.application.errors.pricing import PricingSettingsNotFoundError
from memiro.application.export_pricing_workbook import (
    ExportPricingWorkbook,
    ExportPricingWorkbookForm,
    PricingWorkbookFile,
)
from tests.common.factory.catalog import CATEGORY, PRODUCT
from tests.integration.prime import prime_no_pricing_settings

pytestmark = pytest.mark.usefixtures("catalog")

# What the demo mirror costs at 800 × 600 — the canonical case the workbook
# was checked against (§14.6.7 — hardcoded, never re-derived).
CANONICAL_TOTAL = Decimal(8900)


async def _exported(container: AsyncContainer, form: ExportPricingWorkbookForm) -> PricingWorkbookFile:
    """Run the export in its own production REQUEST scope and take back the file."""
    async with container() as request:
        interactor = await request.get(ExportPricingWorkbook)
        return await interactor.execute(form)


async def test_the_owner_takes_the_workbook_of_a_named_product(container: AsyncContainer) -> None:
    """The book of a product is named after it and opens as a real spreadsheet."""
    exported = await _exported(container, ExportPricingWorkbookForm(product_id=PRODUCT))

    workbook = load_workbook(BytesIO(exported.content))
    assert exported.name == "raschet-zerkalo-v-rame.xlsx"
    assert CHECK in workbook.sheetnames


async def test_the_check_sheet_of_the_workbook_carries_the_engine_total(container: AsyncContainer) -> None:
    """The canonical mirror costs 8 900 ₽, and the book says so with the engine's own number."""
    exported = await _exported(container, ExportPricingWorkbookForm(product_id=PRODUCT))

    sheet = load_workbook(BytesIO(exported.content))[CHECK]
    totals = [sheet.cell(row=row, column=3).value for row in range(DATA_START, DATA_START + 4)]
    assert CANONICAL_TOTAL in totals


async def test_the_owner_takes_the_workbook_of_a_section_without_naming_a_product(
    container: AsyncContainer,
) -> None:
    """A section is enough for a book: it is the dictionary that is being tinkered with."""
    exported = await _exported(container, ExportPricingWorkbookForm(category_id=CATEGORY))

    assert exported.name == "raschet.xlsx"
    assert exported.content[:2] == b"PK"


async def test_the_workbook_of_a_section_has_nothing_to_check(container: AsyncContainer) -> None:
    """Without a product there are no engine totals, and the check sheet says why."""
    exported = await _exported(container, ExportPricingWorkbookForm(category_id=CATEGORY))

    sheet = load_workbook(BytesIO(exported.content))[CHECK]
    assert sheet.cell(row=DATA_START, column=3).value is None


async def test_fails_if_the_named_product_is_not_in_the_catalogue(container: AsyncContainer) -> None:
    """PRODUCT_NOT_FOUND: there is no markup to open the book on."""
    with pytest.raises(ProductNotFoundError):
        await _exported(container, ExportPricingWorkbookForm(product_id=uuid4()))


async def test_fails_if_the_named_section_is_not_in_the_catalogue(container: AsyncContainer) -> None:
    """CATEGORY_NOT_FOUND: there is no dictionary to collect."""
    with pytest.raises(CategoryNotFoundError):
        await _exported(container, ExportPricingWorkbookForm(category_id=uuid4()))


async def test_fails_if_nothing_was_named_at_all() -> None:
    """A request naming neither a product nor a section does not reach the interactor."""
    with pytest.raises(ValueError, match="product or a category"):
        ExportPricingWorkbookForm()


async def test_fails_if_the_site_has_no_pricing_settings(
    container: AsyncContainer,
    engine: AsyncEngine,
) -> None:
    """PRICING_SETTINGS_NOT_FOUND: without the bounds of the calculation the book has nothing to repeat."""
    await prime_no_pricing_settings(engine)

    with pytest.raises(PricingSettingsNotFoundError):
        await _exported(container, ExportPricingWorkbookForm(product_id=PRODUCT))
