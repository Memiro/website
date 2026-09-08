"""The workbook against the engine: does it carry the same dictionary and end the sum the same way."""

from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from memiro.adapters.xlsx.labels import UNIT_LABELS, key_of
from memiro.adapters.xlsx.pricing_workbook import (
    CALCULATION,
    CHECK,
    DATA_START,
    DICTIONARY,
    SETTINGS,
    TABLE_START,
    OpenpyxlPricingWorkbookRenderer,
)
from memiro.application.common.gateway.workbook import CheckedSize, PricingWorkbookSource
from memiro.entities.catalog.attribute.entity import Attribute
from memiro.entities.catalog.attribute.rate import Unit
from memiro.entities.catalog.product.entity import Product
from memiro.entities.common.measure import Dimensions, Millimeters
from memiro.entities.pricing.pricing_service import ROUNDING_STEP, price_product
from memiro.entities.pricing.pricing_settings import PricingSettings
from tests.common.factory.catalog import (
    demo_attributes,
    demo_cutouts,
    demo_numeric_product,
    demo_product,
    demo_settings,
    demo_size_surcharge,
)

CHECKED = ((400, 300), (800, 600), (1200, 700), (2300, 900))

# What the numeric product declares: two cut-outs, hardcoded the way an
# expected value is (§14.6.7).
CUTOUTS_DECLARED = Decimal(2)


def _source(
    *,
    attributes: list[Attribute] | None = None,
    product: Product | None = None,
    settings: PricingSettings | None = None,
    checks: tuple[CheckedSize, ...] = (),
) -> PricingWorkbookSource:
    """Build the source the renderer is given, with the demo catalogue as its default."""
    return PricingWorkbookSource(
        attributes=tuple(attributes if attributes is not None else demo_attributes()),
        settings=settings if settings is not None else demo_settings(size_surcharges=[demo_size_surcharge()]),
        product=product,
        checks=checks,
    )


def _rendered(source: PricingWorkbookSource) -> Workbook:
    """Render the workbook and open it back the way the owner's spreadsheet does."""
    return load_workbook(BytesIO(OpenpyxlPricingWorkbookRenderer().render(source)))


def _checked_by_the_engine(
    product: Product, attributes: list[Attribute], settings: PricingSettings
) -> tuple[CheckedSize, ...]:
    """Price the check sizes with the single implementation of the calculation."""
    sizes = tuple(
        Dimensions(width=Millimeters(value=width), height=Millimeters(value=height)) for width, height in CHECKED
    )
    return tuple(
        CheckedSize(
            dimensions=dimensions,
            quotation=price_product(
                product=product,
                attributes=attributes,
                settings=settings,
                dimensions=dimensions,
                selections={},
            ),
        )
        for dimensions in sizes
    )


def test_the_dictionary_carries_every_value_of_the_section() -> None:
    """The workbook's dictionary is the owner's dictionary, without losses."""
    attributes = demo_attributes()

    sheet = _rendered(_source(attributes=attributes))[DICTIONARY]

    written = {
        (sheet.cell(row=row, column=2).value, sheet.cell(row=row, column=3).value)
        for row in range(DATA_START, sheet.max_row + 1)
    }
    assert written == {(attribute.name, value.name) for attribute in attributes for value in attribute.values}


def test_the_calculation_builds_the_key_the_dictionary_is_written_with() -> None:
    """The lookup key of a row is assembled by the same rule that wrote the dictionary column."""
    workbook = _rendered(_source(product=demo_product()))

    formula = workbook[CALCULATION].cell(row=TABLE_START, column=4).value

    separator = str(formula).split('"')[1]
    assert key_of("Рама", "Алюминий") == f"Рама{separator}Алюминий"


def test_the_units_are_spelled_the_way_the_quantity_formula_reads_them() -> None:
    """A unit word in the dictionary is a word the sum column knows how to charge."""
    workbook = _rendered(_source(product=demo_product()))
    dictionary = workbook[DICTIONARY]

    formula = str(workbook[CALCULATION].cell(row=TABLE_START, column=7).value)

    spelled = {str(dictionary.cell(row=row, column=4).value) for row in range(DATA_START, dictionary.max_row + 1)}
    assert spelled <= set(UNIT_LABELS.values())
    assert all(UNIT_LABELS[unit] in formula for unit in (Unit.SQUARE_METER, Unit.LINEAR_METER, Unit.PIECE))


def test_the_rounding_step_of_the_book_is_the_engine_one() -> None:
    """The last operation of the calculation rounds by the production constant, not by a copy of it."""
    sheet = _rendered(_source(product=demo_product()))[SETTINGS]

    assert sheet["B8"].value == ROUNDING_STEP


def test_the_bounds_of_the_book_are_the_ones_the_owner_set() -> None:
    """The minimum area and the minimum order total arrive from the settings row, unrounded."""
    settings = demo_settings()

    sheet = _rendered(_source(settings=settings))[SETTINGS]

    assert (sheet["B4"].value, sheet["B5"].value) == (settings.min_area.value, settings.min_order_total.amount)


def test_the_markup_of_the_named_product_becomes_the_defaults() -> None:
    """The book opens on what the product declares, not on the first value that came along."""
    attributes = demo_attributes()

    sheet = _rendered(_source(attributes=attributes, product=demo_product()))[CALCULATION]

    defaults = {
        sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
        for row in range(TABLE_START, TABLE_START + len(attributes))
    }
    assert defaults == {
        "Тип полотна": "Серебро",
        "Форма": "Прямоугольное",
        "Рама": "Алюминий",
        "Подсветка": "Без подсветки",
        "Крепление": "С креплением",
        "Подогрев": None,
    }


def test_an_attribute_the_product_never_declared_opens_empty() -> None:
    """A value invented for an undeclared attribute would make the book dearer than the site."""
    attributes = [*demo_attributes(), demo_cutouts()]

    sheet = _rendered(_source(attributes=attributes, product=demo_product()))[CALCULATION]

    cutouts = next(
        row
        for row in range(TABLE_START, TABLE_START + len(attributes))
        if sheet.cell(row=row, column=1).value == "Вырезы"
    )
    assert sheet.cell(row=cutouts, column=2).value is None


def test_a_numeric_attribute_brings_the_quantity_the_product_declared() -> None:
    """A number is a consumption, and the book opens on the one the owner declared."""
    attributes = [*demo_attributes(), demo_cutouts()]

    sheet = _rendered(_source(attributes=attributes, product=demo_numeric_product(quantity=Decimal(2))))[CALCULATION]

    cutouts = next(
        row
        for row in range(TABLE_START, TABLE_START + len(attributes))
        if sheet.cell(row=row, column=1).value == "Вырезы"
    )
    assert sheet.cell(row=cutouts, column=2).value == CUTOUTS_DECLARED


def test_a_numeric_row_is_charged_by_its_own_dictionary_row() -> None:
    """A number has one tariff, so its key is written out instead of being read off the choice."""
    attributes = [*demo_attributes(), demo_cutouts()]

    sheet = _rendered(_source(attributes=attributes, product=demo_numeric_product(quantity=Decimal(2))))[CALCULATION]

    cutouts = next(
        row
        for row in range(TABLE_START, TABLE_START + len(attributes))
        if sheet.cell(row=row, column=1).value == "Вырезы"
    )
    assert sheet.cell(row=cutouts, column=4).value == key_of("Вырезы", "Вырез")


def test_the_check_sheet_shows_exactly_what_the_engine_answered() -> None:
    """The check sheet is the engine's own totals: a book that disagrees with them is the defect."""
    attributes = demo_attributes()
    settings = demo_settings(size_surcharges=[demo_size_surcharge()])
    product = demo_product()
    checks = _checked_by_the_engine(product, attributes, settings)

    sheet = _rendered(_source(attributes=attributes, product=product, settings=settings, checks=checks))[CHECK]

    written = [sheet.cell(row=row, column=3).value for row in range(DATA_START, DATA_START + len(checks))]
    assert written == [
        checked.quotation.total.amount for checked in checks if checked.quotation and checked.quotation.total
    ]


def test_a_book_without_a_product_says_why_it_has_nothing_to_check() -> None:
    """An empty check sheet is an answer, not a hole: the owner asked for a section, not for a mirror."""
    sheet = _rendered(_source())[CHECK]

    assert sheet.cell(row=DATA_START, column=1).value is not None


def test_a_book_of_an_empty_section_is_still_a_book() -> None:
    """A section nobody has priced yet opens without a crash and without a single line."""
    workbook = _rendered(_source(attributes=[]))

    assert workbook[CALCULATION].cell(row=TABLE_START, column=1).value is None


@pytest.mark.parametrize("sheet_name", [CALCULATION, DICTIONARY, SETTINGS, CHECK])
def test_the_book_holds_the_sheet(sheet_name: str) -> None:
    """Every sheet the guide sends the owner to is in the book."""
    workbook = _rendered(_source(product=demo_product()))

    assert sheet_name in workbook.sheetnames
