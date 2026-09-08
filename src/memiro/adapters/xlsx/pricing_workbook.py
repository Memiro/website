from collections.abc import Sequence
from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from typing import override

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from memiro.adapters.xlsx.labels import NO, UNIT_LABELS, YES, key_of
from memiro.application.common.gateway.workbook import PricingWorkbookRenderer, PricingWorkbookSource
from memiro.entities.catalog.attribute.chosen_value import ChosenValue
from memiro.entities.catalog.attribute.entity import Attribute, AttributeKind, AttributeValue
from memiro.entities.catalog.attribute.rate import Unit
from memiro.entities.catalog.product.entity import Product
from memiro.entities.common.identifiers import AttributeId
from memiro.entities.pricing.pricing_service import ROUNDING_STEP
from memiro.entities.pricing.pricing_settings import PricingSettings

GUIDE = "Как пользоваться"
CALCULATION = "Расчёт"
DICTIONARY = "Справочник"
SETTINGS = "Параметры расчёта"
SURCHARGE = "Наценка за размер"
CHECK = "Сверка"

# Every sheet keeps its heading in the first rows and its data below them, so
# the row a lookup range starts at is one number for the whole book.
DATA_START = 5

# The columns of the dictionary, by the position ``VLOOKUP`` asks them by.
UNIT_COLUMN = 4
RATE_COLUMN = 5
SHAPE_COLUMN = 6
SIZE_COLUMN = 7
ABSENCE_COLUMN = 8

INK = "1A2422"
EDGE = "0E6656"
FAINT = "8A9793"

HEAD = PatternFill("solid", fgColor="E2EFEB")
INPUT = PatternFill("solid", fgColor="FFF6E3")
TOTAL = PatternFill("solid", fgColor="F6EDE2")
THIN = Side(style="thin", color="C9D5D1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY = '#,##0.##" ₽"'
MONEY_ROUND = '#,##0" ₽"'
FACTOR = '0.00"×"'
MEASURE = "0.###"

WIDTH_CELL = "$B$3"
HEIGHT_CELL = "$B$4"
AREA_CELL = "$B$5"
PERIMETER_CELL = "$B$6"
LONG_SIDE_CELL = "$B$7"
SHAPE_FACTOR_CELL = "$B$8"
SIZE_FACTOR_CELL = "$B$9"
WITHIN_LIMITS_CELL = "$B$10"
TABLE_HEADER_ROW = 13
TABLE_START = TABLE_HEADER_ROW + 1

MIN_AREA_CELL = f"'{SETTINGS}'!$B$4"
MIN_ORDER_CELL = f"'{SETTINGS}'!$B$5"
MAX_LONG_CELL = f"'{SETTINGS}'!$B$6"
MAX_SHORT_CELL = f"'{SETTINGS}'!$B$7"
ROUNDING_CELL = f"'{SETTINGS}'!$B$8"

SQUARE_METER = UNIT_LABELS[Unit.SQUARE_METER]
LINEAR_METER = UNIT_LABELS[Unit.LINEAR_METER]
PIECE = UNIT_LABELS[Unit.PIECE]
SHAPE_UNIT = UNIT_LABELS[Unit.FACTOR]

NOT_PRICEABLE = "Движок не считает этот товар: у него не заполнены все атрибуты раздела или ни одна статья не платная."
NO_PRODUCT = "Товар не назван — книга собрана по разделу, и сверять её не с чем."
BEYOND_LIMITS = "За производственными границами"


@dataclass(frozen=True, slots=True)
class _Line:
    """One attribute of the calculation sheet with the row it was written to."""

    attribute: Attribute
    row: int


class OpenpyxlPricingWorkbookRenderer(PricingWorkbookRenderer):
    """openpyxl-based implementation of ``PricingWorkbookRenderer``."""

    @override
    def render(self, source: PricingWorkbookSource) -> bytes:
        """Write the six sheets of the workbook and hand back the file."""
        workbook = Workbook()
        # A new book comes with one sheet nobody asked for, and the sheets of
        # this one are all created by name below.
        del workbook[workbook.sheetnames[0]]
        attributes = sorted(source.attributes, key=lambda attribute: (attribute.sort_order, str(attribute.id)))
        _guide_sheet(workbook.create_sheet(GUIDE))
        calculation = workbook.create_sheet(CALCULATION)
        ranges = _dictionary_sheet(workbook.create_sheet(DICTIONARY), attributes)
        _settings_sheet(workbook.create_sheet(SETTINGS), source.settings)
        surcharges = _surcharge_sheet(workbook.create_sheet(SURCHARGE), source.settings)
        _check_sheet(workbook.create_sheet(CHECK), source)
        _calculation_sheet(calculation, attributes, source.product, ranges, surcharges)
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()


@dataclass(frozen=True, slots=True)
class _Ranges:
    """Where the dictionary of the book lives, in the terms its formulas use."""

    lookup: str
    values: dict[AttributeId, str]


def _title(sheet: Worksheet, cell: str, text: str, size: int = 16) -> None:
    """Put the heading of a sheet in its first cell."""
    sheet[cell] = text
    sheet[cell].font = Font(bold=True, size=size, color=INK)


def _note(sheet: Worksheet, cell: str, text: str) -> None:
    """Put a faint line of explanation under a heading."""
    sheet[cell] = text
    sheet[cell].font = Font(size=10, italic=True, color=FAINT)


def _header(sheet: Worksheet, row: int, titles: Sequence[str]) -> None:
    """Lay out the header row of a table."""
    for column, text in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=column, value=text)
        cell.font = Font(bold=True, size=10, color=EDGE)
        cell.fill = HEAD
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _widths(sheet: Worksheet, sizes: dict[str, int]) -> None:
    """Set the column widths of a sheet."""
    for column, width in sizes.items():
        sheet.column_dimensions[column].width = width


def _label(sheet: Worksheet, row: int, text: str) -> None:
    """Write the name of one parameter in the left column of a block."""
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = Font(size=11, color=INK)


def _input(sheet: Worksheet, row: int, value: str | int | Decimal, number_format: str) -> None:
    """Write a cell the owner is meant to type in."""
    cell = sheet.cell(row=row, column=2, value=value)
    cell.fill = INPUT
    cell.border = BOX
    cell.number_format = number_format


def _derived(sheet: Worksheet, row: int, value: str | int | Decimal, number_format: str) -> None:
    """Write a cell the book computes for itself, or a bound it was handed."""
    cell = sheet.cell(row=row, column=2, value=value)
    cell.number_format = number_format
    cell.font = Font(size=11, color=INK)


def _yes_no(*, flag: bool) -> str:
    """Spell a domain flag the way the formulas of the book read it."""
    return YES if flag else NO


def _dictionary_sheet(sheet: Worksheet, attributes: Sequence[Attribute]) -> _Ranges:
    """Write the whole dictionary of the section — every value, its tariff and its flags."""
    _title(sheet, "A1", DICTIONARY)
    _note(sheet, "A2", "Все цены раздела. Единица расхода решает, на что умножается тариф.")
    _widths(sheet, {"A": 42, "B": 24, "C": 26, "D": 14, "E": 14, "F": 16, "G": 18, "H": 14})
    _header(
        sheet,
        DATA_START - 1,
        (
            "Ключ",
            "Атрибут",
            "Значение",
            "Единица расхода",
            "Тариф",
            "Умножается формой",
            "Умножается наценкой за размер",
            "Означает отсутствие",
        ),
    )
    row = DATA_START
    values: dict[AttributeId, str] = {}
    for attribute in attributes:
        first = row
        for value in _ordered(attribute.values):
            _dictionary_row(sheet, row, attribute, value)
            row += 1
        values[attribute.id] = f"'{DICTIONARY}'!$C${first}:$C${row - 1}"
    last = max(row - 1, DATA_START)
    return _Ranges(lookup=f"'{DICTIONARY}'!$A${DATA_START}:$H${last}", values=values)


def _ordered(values: Sequence[AttributeValue]) -> list[AttributeValue]:
    """Put the values of an attribute in the owner's order, breaking a tie the way the engine does."""
    return sorted(values, key=lambda value: (value.sort_order, str(value.id)))


def _dictionary_row(sheet: Worksheet, row: int, attribute: Attribute, value: AttributeValue) -> None:
    """Write one dictionary value with everything the calculation looks up about it."""
    cells: tuple[str | Decimal, ...] = (
        key_of(attribute.name, value.name),
        attribute.name,
        value.name,
        UNIT_LABELS[value.rate.unit],
        value.rate.amount.amount,
        _yes_no(flag=value.scaled_by_shape),
        _yes_no(flag=value.scaled_by_size_surcharge),
        _yes_no(flag=value.marks_absence),
    )
    for column, content in enumerate(cells, start=1):
        cell = sheet.cell(row=row, column=column, value=content)
        cell.border = BOX
        cell.font = Font(size=10, color=INK)
    sheet.cell(row=row, column=RATE_COLUMN + 1).number_format = FACTOR if value.rate.unit is Unit.FACTOR else MONEY


def _settings_sheet(sheet: Worksheet, settings: PricingSettings) -> None:
    """Write the bounds of the calculation the formulas of the book lean on."""
    _title(sheet, "A1", SETTINGS)
    _note(sheet, "A2", "Границы расчёта из админки. Ноль в максимальной стороне значит «без границы».")
    _widths(sheet, {"A": 42, "B": 18})
    bounds: tuple[tuple[str, int | Decimal, str], ...] = (
        ("Минимальная площадь расчёта, м²", settings.min_area.value, MEASURE),
        ("Минимальная сумма заказа", settings.min_order_total.amount, MONEY),
        ("Максимальная длинная сторона, мм", settings.max_long_side_mm.value, "0"),
        ("Максимальная короткая сторона, мм", settings.max_short_side_mm.value, "0"),
        ("Шаг округления итога", ROUNDING_STEP, MONEY_ROUND),
    )
    for row, (label, value, number_format) in enumerate(bounds, start=4):
        _label(sheet, row, label)
        _derived(sheet, row, value, number_format)


def _surcharge_sheet(sheet: Worksheet, settings: PricingSettings) -> str | None:
    """Write the size-surcharge tiers and tell the calculation where to look them up."""
    _title(sheet, "A1", SURCHARGE)
    _note(sheet, "A2", "Ступени для крупных изделий: множитель берёт та, чей порог перешагнула длинная сторона.")
    _widths(sheet, {"A": 32, "B": 16})
    _header(sheet, DATA_START - 1, ("От длинной стороны, мм", "Множитель"))
    tiers = sorted(settings.size_surcharges, key=lambda tier: tier.from_long_side_mm.value)
    for row, tier in enumerate(tiers, start=DATA_START):
        threshold = sheet.cell(row=row, column=1, value=tier.from_long_side_mm.value)
        threshold.border = BOX
        factor = sheet.cell(row=row, column=2, value=tier.factor)
        factor.border = BOX
        factor.number_format = FACTOR
    if not tiers:
        _note(sheet, f"A{DATA_START}", "Ступеней нет: крупный размер стоит столько же, сколько обычный.")
        return None
    last = DATA_START + len(tiers) - 1
    return f"'{SURCHARGE}'!$A${DATA_START}:$A${last},'{SURCHARGE}'!$B${DATA_START}:$B${last}"


def _calculation_sheet(
    sheet: Worksheet,
    attributes: Sequence[Attribute],
    product: Product | None,
    ranges: _Ranges,
    surcharges: str | None,
) -> None:
    """Write the calculator itself: sizes on top, the configuration below, the total at the bottom."""
    _title(sheet, "A1", CALCULATION, 18)
    _widths(sheet, {"A": 30, "B": 26, "C": 26, "D": 40, "E": 14, "F": 14, "G": 14, "H": 12, "I": 14, "J": 16})
    _sizes_block(sheet, surcharges)
    _note(
        sheet,
        "A12",
        "Жёлтые ячейки вводятся руками, остальные — формулы. «Ключ» и «Учитывается» служебные: "
        "по ним книга находит тариф.",
    )
    _header(
        sheet,
        TABLE_HEADER_ROW,
        (
            "Атрибут",
            "Умолчание товара",
            "Выбор покупателя",
            "Ключ",
            "Единица расхода",
            "Тариф",
            "Количество",
            "Вклад в форму",
            "Учитывается",
            "Сумма",
        ),
    )
    lines = _lines(sheet, attributes, product, ranges)
    for line in lines:
        _line_formulas(sheet, line, lines, ranges)
    _totals_block(sheet, lines)


def _sizes_block(sheet: Worksheet, surcharges: str | None) -> None:
    """Write the sizes the owner types and everything the book derives from them."""
    _label(sheet, 3, "Ширина, мм")
    _input(sheet, 3, 800, "0")
    _label(sheet, 4, "Высота, мм")
    _input(sheet, 4, 600, "0")
    _label(sheet, 5, "Площадь, м²")
    _derived(sheet, 5, f"=MAX({WIDTH_CELL}*{HEIGHT_CELL}/1000000,{MIN_AREA_CELL})", MEASURE)
    _label(sheet, 6, "Периметр, пог. м")
    _derived(sheet, 6, f"=2*({WIDTH_CELL}+{HEIGHT_CELL})/1000", MEASURE)
    _label(sheet, 7, "Длинная сторона, мм")
    _derived(sheet, 7, f"=MAX({WIDTH_CELL},{HEIGHT_CELL})", "0")
    _label(sheet, 8, "Коэффициент формы")
    _derived(sheet, 8, "=1", FACTOR)
    _label(sheet, 9, "Наценка за размер")
    _derived(
        sheet,
        9,
        f"=IFERROR(LOOKUP({LONG_SIDE_CELL},{surcharges}),1)" if surcharges is not None else "=1",
        FACTOR,
    )
    _label(sheet, 10, "В производственных границах")
    _derived(
        sheet,
        10,
        f"=IF(AND(OR({MAX_LONG_CELL}=0,{LONG_SIDE_CELL}<={MAX_LONG_CELL}),"
        f"OR({MAX_SHORT_CELL}=0,MIN({WIDTH_CELL},{HEIGHT_CELL})<={MAX_SHORT_CELL})),"
        f'"{YES}","{NO}")',
        "General",
    )


def _lines(
    sheet: Worksheet,
    attributes: Sequence[Attribute],
    product: Product | None,
    ranges: _Ranges,
) -> list[_Line]:
    """Write one row per attribute: its name, what the product declared and what is chosen now."""
    declared = (
        {declaration.attribute_id: declaration.chosen for declaration in product.declared_values} if product else {}
    )
    lines: list[_Line] = []
    for row, attribute in enumerate(attributes, start=TABLE_START):
        values = _ordered(attribute.values)
        default = _default_of(attribute, values, declared.get(attribute.id), named=product is not None)
        name = sheet.cell(row=row, column=1, value=attribute.name)
        name.font = Font(size=11, color=INK)
        name.border = BOX
        stated = sheet.cell(row=row, column=2, value=default)
        stated.font = Font(size=10, color=FAINT)
        stated.border = BOX
        choice = sheet.cell(row=row, column=3, value=default)
        choice.fill = INPUT
        choice.border = BOX
        if attribute.kind is AttributeKind.SELECT and values:
            _dropdown(sheet, choice.coordinate, ranges.values[attribute.id])
        lines.append(_Line(attribute=attribute, row=row))
    return lines


def _default_of(
    attribute: Attribute,
    values: Sequence[AttributeValue],
    chosen: ChosenValue | None,
    *,
    named: bool,
) -> str | int | Decimal:
    """Say what the row opens on: the product's declaration, or the first value when no product was named.

    An attribute the named product never declared opens empty on purpose: the
    engine prices no such attribute either, and a value invented here would
    make the book dearer than the site.
    """
    if chosen is not None:
        if attribute.kind is AttributeKind.NUMBER:
            return chosen.quantity if chosen.quantity is not None else 0
        declared = next((value for value in values if value.id == chosen.value_id), None)
        return declared.name if declared is not None else ""
    if named or not values:
        return ""
    return 0 if attribute.kind is AttributeKind.NUMBER else values[0].name


def _dropdown(sheet: Worksheet, coordinate: str, values: str) -> None:
    """Offer one cell the values of its own attribute and nothing else."""
    validation = DataValidation(type="list", formula1=values, allow_blank=True, showDropDown=False)
    sheet.add_data_validation(validation)
    validation.add(coordinate)  # pyright: ignore[reportUnknownMemberType]


def _line_formulas(sheet: Worksheet, line: _Line, lines: Sequence[_Line], ranges: _Ranges) -> None:
    """Charge one row the way the engine charges the value it names."""
    row = line.row
    numeric = line.attribute.kind is AttributeKind.NUMBER
    key = (
        key_of(line.attribute.name, _ordered(line.attribute.values)[0].name)
        if numeric and line.attribute.values
        else f'=$A{row}&" | "&$C{row}'
    )
    sheet.cell(row=row, column=4, value=key).font = Font(size=9, color=FAINT)
    sheet.cell(row=row, column=5, value=f'=IFERROR(VLOOKUP($D{row},{ranges.lookup},{UNIT_COLUMN},FALSE),"")')
    rate = sheet.cell(row=row, column=6, value=f"=IFERROR(VLOOKUP($D{row},{ranges.lookup},{RATE_COLUMN},FALSE),0)")
    rate.number_format = MONEY
    quantity = sheet.cell(
        row=row,
        column=7,
        value=f"=$C{row}"
        if numeric
        else f'=IF($E{row}="{SQUARE_METER}",{AREA_CELL},'
        f'IF($E{row}="{LINEAR_METER}",{PERIMETER_CELL},'
        f'IF($E{row}="{PIECE}",1,0)))',
    )
    quantity.number_format = MEASURE
    shape = sheet.cell(
        row=row,
        column=8,
        value=f'=IF(AND($I{row}="{YES}",$E{row}="{SHAPE_UNIT}"),$F{row},1)',
    )
    shape.number_format = FACTOR
    sheet.cell(row=row, column=9, value=_applies_formula(line, lines, ranges)).font = Font(size=9, color=FAINT)
    amount = sheet.cell(
        row=row,
        column=10,
        value=f'=IF(OR($I{row}="{NO}",$E{row}="{SHAPE_UNIT}"),0,'
        f"$F{row}*$G{row}"
        f'*IF(VLOOKUP($D{row},{ranges.lookup},{SHAPE_COLUMN},FALSE)="{YES}",{SHAPE_FACTOR_CELL},1)'
        f'*IF(VLOOKUP($D{row},{ranges.lookup},{SIZE_COLUMN},FALSE)="{YES}",{SIZE_FACTOR_CELL},1))',
    )
    amount.number_format = MONEY
    amount.border = BOX


def _applies_formula(line: _Line, lines: Sequence[_Line], ranges: _Ranges) -> str:
    """Say whether a dependent attribute is applicable — the rule the engine applies to its parents."""
    rows = {other.attribute.id: other.row for other in lines}
    parents = [rows[parent_id] for parent_id in line.attribute.parent_ids if parent_id in rows]
    if not parents:
        return YES
    present = ",".join(
        f'AND($I{parent}="{YES}",IFERROR(VLOOKUP($D{parent},{ranges.lookup},{ABSENCE_COLUMN},FALSE),"{NO}")<>"{YES}")'
        for parent in parents
    )
    return f'=IF(OR({present}),"{YES}","{NO}")'


def _totals_block(sheet: Worksheet, lines: Sequence[_Line]) -> None:
    """Add the lines up the way the engine ends the calculation: threshold, then rounding."""
    last = lines[-1].row if lines else TABLE_START
    amounts = f"J{TABLE_START}:J{last}"
    row = last + 2
    if lines:
        sheet[SHAPE_FACTOR_CELL] = f"=PRODUCT(H{TABLE_START}:H{last})"
    _label(sheet, row, "Сумма статей")
    _derived(sheet, row, f"=SUM({amounts})" if lines else "=0", MONEY)
    _label(sheet, row + 1, "После минимальной суммы заказа")
    _derived(sheet, row + 1, f"=MAX(B{row},{MIN_ORDER_CELL})", MONEY)
    _label(sheet, row + 2, "Итог")
    total = sheet.cell(
        row=row + 2,
        column=2,
        value=f'=IF({WITHIN_LIMITS_CELL}="{NO}","{BEYOND_LIMITS}",'
        f"ROUNDUP(B{row + 1}/{ROUNDING_CELL},0)*{ROUNDING_CELL})",
    )
    total.number_format = MONEY_ROUND
    total.fill = TOTAL
    total.border = BOX
    total.font = Font(bold=True, size=12, color=INK)


def _check_sheet(sheet: Worksheet, source: PricingWorkbookSource) -> None:
    """Write what the engine itself answers on a handful of sizes — the reason the book can be trusted."""
    _title(sheet, "A1", CHECK)
    _note(sheet, "A2", "Итоги посчитал движок сайта. Соберите то же зеркало на «Расчёте» — числа обязаны совпасть.")
    _widths(sheet, {"A": 16, "B": 16, "C": 20, "D": 26, "E": 46})
    _header(sheet, DATA_START - 1, ("Ширина, мм", "Высота, мм", "Итог движка", "Наценка за размер, от мм", "Статей"))
    if not source.checks:
        _note(sheet, f"A{DATA_START}", NO_PRODUCT if source.product is None else NOT_PRICEABLE)
        return
    for row, checked in enumerate(source.checks, start=DATA_START):
        sheet.cell(row=row, column=1, value=checked.dimensions.width.value).border = BOX
        sheet.cell(row=row, column=2, value=checked.dimensions.height.value).border = BOX
        quotation = checked.quotation
        total = sheet.cell(
            row=row, column=3, value=quotation.total.amount if quotation and quotation.total else BEYOND_LIMITS
        )
        total.number_format = MONEY_ROUND
        total.border = BOX
        threshold = quotation.size_surcharge_from_long_side_mm if quotation is not None else None
        sheet.cell(row=row, column=4, value=threshold.value if threshold is not None else "—").border = BOX
        sheet.cell(row=row, column=5, value=len(quotation.breakdown) if quotation is not None else 0).border = BOX


def _guide_sheet(sheet: Worksheet) -> None:
    """Write the first sheet: what the book is and where to press."""
    _title(sheet, "A1", "Как пользоваться книгой", 18)
    sheet.sheet_view.showGridLines = False
    _widths(sheet, {"A": 4, "B": 30, "C": 100})
    blocks = (
        (
            "Что это",
            (
                "Точная копия расчёта, который сайт делает на карточке товара, собранная из живых данных админки. "
                "Крутите её сколько угодно: база от этого не меняется."
            ),
        ),
        (
            "Жёлтые ячейки",
            "Всё, что вводится руками: размеры и выбор покупателя. Остальное — формулы, их лучше не трогать.",
        ),
        ("", ""),
        (
            "Лист «Расчёт»",
            (
                "Собираете зеркало: размеры сверху, характеристики выпадающими списками. «Умолчание товара» — то, "
                "чем размечен товар, «Выбор покупателя» — то, что крутит он."
            ),
        ),
        (
            "Лист «Справочник»",
            "Все цены раздела. Единица расхода решает, на что умножается тариф: на площадь, на периметр или на штуку.",
        ),
        (
            "Лист «Наценка за размер»",
            "Ступени для крупных изделий: множитель берёт та, чей порог перешагнула длинная сторона.",
        ),
        (
            "Лист «Сверка»",
            (
                "Итоги, посчитанные движком сайта. Соберите то же зеркало на «Расчёте» — числа обязаны совпасть, "
                "а расхождение стоит показать разработчику."
            ),
        ),
        ("", ""),
        (
            "Три единицы расхода",
            (
                "За квадратный метр — полотно: расходуется листом. За погонный метр — рама, кромка, лента: идут по "
                "контуру. За штуку — кнопка, подогрев, крепление, вырез."
            ),
        ),
        (
            "Коэффициент формы",
            (
                "Умножает только то, что помечено в справочнике колонкой «умножается формой»: на круге дороже режется "
                "полотно, а лента не дорожает."
            ),
        ),
        (
            "Порядок сложения",
            (
                "Статьи складываются, итог поднимается до минимальной суммы заказа и округляется вверх до шага "
                "округления — ровно в этом порядке."
            ),
        ),
        (
            "Чего в книге нет",
            (
                "Доплаты покупателя за отдельный выбор: её сайт считает по каждому атрибуту отдельно. Книга отвечает "
                "на другой вопрос — сколько стоит собранное зеркало целиком."
            ),
        ),
    )
    row = 3
    for head, text in blocks:
        if not head:
            row += 1
            continue
        title = sheet.cell(row=row, column=2, value=head)
        title.font = Font(bold=True, size=11, color=EDGE)
        title.alignment = Alignment(vertical="top")
        body = sheet.cell(row=row, column=3, value=text)
        body.font = Font(size=11, color=INK)
        body.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 15 * (1 + len(text) // 95)
        row += 1
